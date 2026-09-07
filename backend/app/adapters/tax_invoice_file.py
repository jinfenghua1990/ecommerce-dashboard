"""读取税务系统官方导出的发票清单。

这里只做保守的表头识别和字段抽取。原始行由导入服务完整保存；没有明确表头或
发票号码的行进入待核对，不把缺失数据推断成“未开票”。
"""

from __future__ import annotations

import csv
import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable


ALLOWED_EXTENSIONS = {".xlsx", ".csv"}
MAX_XLSX_UNCOMPRESSED_BYTES = 300 * 1024 * 1024

FIELD_ALIASES: dict[str, tuple[str, ...]] = {
    "invoice_number": ("发票号码", "发票号", "发票编号", "数电发票号码", "电子发票号码", "invoice number", "invoice_no"),
    "invoice_code": ("发票代码", "票代码", "数电发票代码", "电子发票代码", "invoice code", "invoice_code"),
    "invoice_type": ("发票类型", "发票种类", "票种", "发票票种", "invoice type"),
    "issue_date": ("开票日期", "开具日期", "开票时间", "发票日期", "issue date", "invoice date"),
    "seller_name": ("销售方名称", "销方名称", "开票方名称", "卖方名称", "seller name"),
    "seller_tax_id": ("销售方识别号", "销方识别号", "开票方税号", "卖方税号", "销售方税号", "seller tax id"),
    "buyer_name": ("购买方名称", "购方名称", "受票方名称", "买方名称", "buyer name"),
    "buyer_tax_id": ("购买方识别号", "购方识别号", "受票方税号", "买方税号", "购买方税号", "buyer tax id"),
    "amount_excl_tax": ("不含税金额", "金额(不含税)", "金额（不含税）", "金额", "amount excl tax", "amount"),
    "tax_amount": ("税额", "税金", "合计税额", "tax amount", "tax"),
    "total_amount": ("价税合计", "价税合计金额", "含税金额", "金额(含税)", "金额（含税）", "发票金额", "total amount"),
    "currency": ("币种", "货币", "currency"),
    "direction": ("发票方向", "进销项", "发票属性", "发票来源", "invoice direction"),
    "status": ("发票状态", "有效状态", "作废标志", "是否作废", "状态", "invoice status"),
    "related_order_ref": ("关联订单号", "订单号", "订单编号", "采购单号", "采购订单号", "业务单号", "来源单号", "原始订单号"),
}


@dataclass(frozen=True)
class ParsedTaxInvoiceExport:
    sheet_name: str
    headers: list[str]
    rows: list[dict[str, str]]
    mapping: dict[str, str]
    direction_hint: str


def normalize_header(value: object) -> str:
    text = unicodedata.normalize("NFKC", str(value or "")).strip().lower()
    return re.sub(r"[\s_\-()（）\[\]【】/\\:：]+", "", text)


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, Decimal):
        return format(value, "f")
    return str(value).strip()


def resolve_mapping(headers: list[str]) -> dict[str, str]:
    normalized = [(header, normalize_header(header)) for header in headers if header]
    mapping: dict[str, str] = {}
    for field, aliases in FIELD_ALIASES.items():
        alias_keys = [normalize_header(alias) for alias in aliases]
        if field == "invoice_number":
            # 数电发票同时可能保留一个空的“发票号码”列；优先使用真正有值的
            # 数电号码列，服务层仍会兼容旧批次的原始映射。
            for preferred in ("数电发票号码", "电子发票号码", "发票号码", "发票号", "发票编号"):
                preferred_key = normalize_header(preferred)
                exact_preferred = next(
                    (header for header, key in normalized if key == preferred_key), None
                )
                if exact_preferred:
                    mapping[field] = exact_preferred
                    break
            if field in mapping:
                continue
        # 先精确命中，避免“金额”之类的短别名抢到相邻字段。
        exact = next((header for header, key in normalized if key in alias_keys), None)
        if exact:
            mapping[field] = exact
            continue
        for alias_key in alias_keys:
            if len(alias_key) < 3:
                continue
            partial = next((header for header, key in normalized if alias_key in key), None)
            if partial:
                mapping[field] = partial
                break
    return mapping


def _header_score(values: list[object]) -> int:
    headers = [cell_text(value) for value in values]
    normalized = [normalize_header(header) for header in headers if header]
    if not normalized:
        return -1
    tokens = ("发票", "开票", "税额", "价税", "税号", "购方", "销方", "金额", "红字", "作废", "票种")
    score = sum(any(token in header for token in tokens) for header in normalized) * 10
    mapping = resolve_mapping(headers)
    score += len(mapping) * 2
    return score + min(len(normalized), 20)


def _find_header_row(table: list[list[object]]) -> int:
    candidates = table[:20]
    if not candidates:
        raise ValueError("文件没有可读取的内容")
    index = max(range(len(candidates)), key=lambda item: _header_score(candidates[item]))
    if _header_score(candidates[index]) < 5:
        raise ValueError("未找到可识别的发票表头")
    return index


def _unique_headers(values: list[object]) -> list[str]:
    result: list[str] = []
    used: dict[str, int] = {}
    for index, value in enumerate(values, start=1):
        base = cell_text(value) or f"列{index}"
        count = used.get(base, 0) + 1
        used[base] = count
        result.append(base if count == 1 else f"{base}#{count}")
    return result


def _rows_to_records(
    rows_iter: Iterable[Iterable[object]], *, header_index: int, max_rows: int
) -> tuple[list[str], list[dict[str, str]]]:
    headers: list[str] | None = None
    rows: list[dict[str, str]] = []
    for source_index, source_row in enumerate(rows_iter):
        if source_index < header_index:
            continue
        raw_row = list(source_row)
        if source_index == header_index:
            headers = _unique_headers(raw_row)
            continue
        if headers is None:
            raise ValueError("未找到有效表头")
        payload: dict[str, str] = {}
        for index, raw_value in enumerate(raw_row[:len(headers)]):
            value = cell_text(raw_value)
            if value:
                payload[headers[index]] = value
        if not payload:
            continue
        if len(rows) >= max_rows:
            raise ValueError(f"数据行超过单次导入上限（{max_rows} 行）")
        rows.append(payload)
    if headers is None:
        raise ValueError("未找到有效表头")
    return headers, rows


def _validate_xlsx(content: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            total = sum(member.file_size for member in archive.infolist())
    except zipfile.BadZipFile as exc:
        raise ValueError("文件不是有效的 XLSX") from exc
    if total > MAX_XLSX_UNCOMPRESSED_BYTES:
        raise ValueError("XLSX 解压后的内容超过安全上限")


def _parse_xlsx(content: bytes, *, max_rows: int) -> tuple[str, list[str], list[dict[str, str]]]:
    from openpyxl import load_workbook

    _validate_xlsx(content)
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("无法读取 XLSX 文件") from exc
    best: tuple[int, str, list[list[object]]] | None = None
    try:
        for worksheet in workbook.worksheets:
            sample = [list(row) for row in worksheet.iter_rows(max_row=20, values_only=True)]
            score = max((_header_score(row) for row in sample), default=-1)
            if best is None or score > best[0]:
                best = (score, worksheet.title, sample)
        if best is None or best[0] < 5:
            raise ValueError("未找到含发票表头的工作表")
        _, sheet_name, sample = best
        worksheet = workbook[sheet_name]
        headers, rows = _rows_to_records(
            worksheet.iter_rows(values_only=True),
            header_index=_find_header_row(sample),
            max_rows=max_rows,
        )
    finally:
        workbook.close()
    return sheet_name, headers, rows


def _decode_csv(content: bytes) -> str:
    for encoding in ("utf-8-sig", "gb18030", "utf-8"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("CSV 编码无法识别，请导出 UTF-8 或 GBK/GB18030 文件")


def _parse_csv(content: bytes, *, max_rows: int) -> tuple[str, list[str], list[dict[str, str]]]:
    text = _decode_csv(content)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",\t;")
    except csv.Error:
        dialect = csv.excel
    table = [list(row) for row in csv.reader(io.StringIO(text), dialect=dialect)]
    header_index = _find_header_row(table)
    headers, rows = _rows_to_records(table, header_index=header_index, max_rows=max_rows)
    return "CSV", headers, rows


def parse_tax_invoice_export(content: bytes, original_name: str, *, max_rows: int) -> ParsedTaxInvoiceExport:
    if not content:
        raise ValueError("空文件")
    extension = Path(original_name or "").suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise ValueError("仅支持税务系统导出的 XLSX 或 CSV 文件")
    if extension == ".xlsx":
        sheet_name, headers, rows = _parse_xlsx(content, max_rows=max_rows)
    else:
        sheet_name, headers, rows = _parse_csv(content, max_rows=max_rows)
    mapping = resolve_mapping(headers)
    normalized_headers = " ".join(normalize_header(value) for value in headers)
    # “进销项”是税务系统常见的综合字段名，不能因为包含“销项”就误判为销项。
    if "进销项" in normalized_headers:
        direction_hint = "unknown"
    elif "进项" in normalized_headers:
        direction_hint = "input"
    elif "销项" in normalized_headers:
        direction_hint = "output"
    else:
        direction_hint = "unknown"
    return ParsedTaxInvoiceExport(sheet_name, headers, rows, mapping, direction_hint)
