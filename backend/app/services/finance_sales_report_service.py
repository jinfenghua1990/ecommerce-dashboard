"""可配置月度销售汇总。

原则：
- 字段、别名、顺序保存在数据库模板中；后台月度任务与前端共用同一模板。
- 有效销售订单口径与 profit.compute 对齐：订单已支付/完成，或支付状态已成功。
- 销售净额按 SalesOrderItem.amount - discount_amount，避免另造利润口径。
- 退款使用本地 AftersalesOrder 已记录的有效退款；取消/关闭/拒绝/作废的售后不计。
- 订单级金额和退款只在该订单第一条明细行输出，避免 Excel 汇总时重复累计。
"""
from __future__ import annotations

from collections import defaultdict
from datetime import datetime, timezone
from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import or_
from sqlalchemy.orm import Session

from app.config import settings
from app.models.finance import ArchiveFile, FinanceSalesReportTemplate
from app.models.sales import AftersalesOrder, SalesOrder, SalesOrderItem
from app.services import finance_service
from app.services.monthly_core import month_bounds
from app.utils.money import quantize

FIELD_REGISTRY: dict[str, dict[str, str]] = {
    "platform": {"label": "平台", "type": "text"},
    "store_id": {"label": "店铺ID", "type": "text"},
    "source_provider": {"label": "数据来源", "type": "text"},
    "order_no": {"label": "订单号", "type": "text"},
    "ordered_at": {"label": "下单时间", "type": "datetime"},
    "paid_at": {"label": "支付时间", "type": "datetime"},
    "order_type": {"label": "订单类型", "type": "text"},
    "order_status": {"label": "订单状态", "type": "text"},
    "pay_status": {"label": "支付状态", "type": "text"},
    "sku_code": {"label": "SKU编码", "type": "text"},
    "goods_name": {"label": "商品名称", "type": "text"},
    "quantity": {"label": "销售数量", "type": "number"},
    "unit_price": {"label": "单价", "type": "money"},
    "item_amount": {"label": "商品金额", "type": "money"},
    "discount_amount": {"label": "优惠金额", "type": "money"},
    "item_net_amount": {"label": "商品净销售额", "type": "money"},
    "order_amount": {"label": "订单金额", "type": "money"},
    "paid_amount": {"label": "实付金额", "type": "money"},
    "refund_amount": {"label": "有效退款金额", "type": "money"},
    "currency": {"label": "币种", "type": "text"},
}

DEFAULT_FIELD_KEYS = [
    "platform",
    "order_no",
    "ordered_at",
    "sku_code",
    "goods_name",
    "quantity",
    "unit_price",
    "item_amount",
    "discount_amount",
    "item_net_amount",
    "paid_amount",
    "refund_amount",
]

DEFAULT_RULES: dict[str, Any] = {
    "valid_order_mode": "paid_or_completed",
    "refund_mode": "recorded_non_cancelled",
    "timezone": settings.TZ,
    "order_level_value_mode": "first_item_only",
}

INVALID_REFUND_STATUS_WORDS = (
    "取消", "关闭", "拒绝", "驳回", "作废", "cancel", "closed", "reject", "void"
)


def default_fields() -> list[dict[str, Any]]:
    return [
        {
            "key": key,
            "label": meta["label"],
            "enabled": key in DEFAULT_FIELD_KEYS,
        }
        for key, meta in FIELD_REGISTRY.items()
    ]


def _normalize_fields(value: Any) -> list[dict[str, Any]]:
    """只接收注册字段；保存数组顺序即导出顺序，同时自动补齐新版本新增字段。"""
    seen: set[str] = set()
    result: list[dict[str, Any]] = []
    for raw in value if isinstance(value, list) else []:
        if not isinstance(raw, dict):
            continue
        key = str(raw.get("key") or "").strip()
        if key not in FIELD_REGISTRY or key in seen:
            continue
        label = str(raw.get("label") or FIELD_REGISTRY[key]["label"]).strip()[:80]
        result.append({"key": key, "label": label or FIELD_REGISTRY[key]["label"],
                       "enabled": bool(raw.get("enabled", True))})
        seen.add(key)
    for key, meta in FIELD_REGISTRY.items():
        if key not in seen:
            result.append({"key": key, "label": meta["label"], "enabled": False})
    return result or default_fields()


def _normalize_emails(value: Any) -> list[str]:
    result: list[str] = []
    for raw in value if isinstance(value, list) else []:
        email = str(raw).strip()
        if email and email not in result:
            result.append(email[:320])
    return result


def get_or_create_template(db: Session, company: str = finance_service.DEFAULT_COMPANY) -> FinanceSalesReportTemplate:
    row = (
        db.query(FinanceSalesReportTemplate)
        .filter_by(company=company, name="默认财务月报")
        .first()
    )
    if row is None:
        row = FinanceSalesReportTemplate(
            company=company,
            name="默认财务月报",
            enabled=True,
            fields=default_fields(),
            rules=dict(DEFAULT_RULES),
            to_addrs=[],
            cc_addrs=[],
            auto_send=False,
            send_day=3,
            send_hour=10,
        )
        db.add(row)
        db.commit()
        db.refresh(row)
    return row


def serialize_template(row: FinanceSalesReportTemplate) -> dict[str, Any]:
    return {
        "id": row.id,
        "company": row.company,
        "name": row.name,
        "enabled": row.enabled,
        "fields": _normalize_fields(row.fields),
        "fieldRegistry": [
            {"key": key, "defaultLabel": meta["label"], "type": meta["type"]}
            for key, meta in FIELD_REGISTRY.items()
        ],
        "rules": {**DEFAULT_RULES, **(row.rules or {})},
        "toAddrs": list(row.to_addrs or []),
        "ccAddrs": list(row.cc_addrs or []),
        "autoSend": row.auto_send,
        "sendDay": row.send_day,
        "sendHour": row.send_hour,
    }


def save_template(
    db: Session,
    *,
    company: str,
    fields: list[dict[str, Any]],
    rules: dict[str, Any] | None,
    to_addrs: list[str] | None,
    cc_addrs: list[str] | None,
    auto_send: bool,
    send_day: int,
    send_hour: int,
    enabled: bool = True,
) -> FinanceSalesReportTemplate:
    if not 1 <= int(send_day) <= 28:
        raise ValueError("自动发送日期只能设置为每月 1-28 日")
    if not 0 <= int(send_hour) <= 23:
        raise ValueError("自动发送小时必须为 0-23")
    row = get_or_create_template(db, company)
    row.fields = _normalize_fields(fields)
    # 当前口径值是受控枚举；前端可调整时仍避免写入未知计算逻辑。
    normalized_rules = dict(DEFAULT_RULES)
    supplied = rules or {}
    if supplied.get("valid_order_mode") in {"paid_or_completed"}:
        normalized_rules["valid_order_mode"] = supplied["valid_order_mode"]
    if supplied.get("refund_mode") in {"recorded_non_cancelled", "ignore_refund"}:
        normalized_rules["refund_mode"] = supplied["refund_mode"]
    normalized_rules["timezone"] = settings.TZ
    normalized_rules["order_level_value_mode"] = "first_item_only"
    row.rules = normalized_rules
    row.to_addrs = _normalize_emails(to_addrs or [])
    row.cc_addrs = _normalize_emails(cc_addrs or [])
    row.auto_send = bool(auto_send)
    row.send_day = int(send_day)
    row.send_hour = int(send_hour)
    row.enabled = bool(enabled)
    db.commit()
    db.refresh(row)
    return row


def _valid_order_condition():
    """与 app.services.profit.compute 的有效销售口径保持一致。"""
    return or_(
        SalesOrder.order_status.in_(["paid", "done", "finished", "已支付", "已完成"]),
        SalesOrder.pay_status.in_(["paid", "success", "已支付"]),
    )


def _is_valid_refund(row: AftersalesOrder) -> bool:
    status = (row.status or "").strip().lower()
    return not any(word in status for word in INVALID_REFUND_STATUS_WORDS)


def _money(value: Decimal | None) -> Decimal:
    return value if value is not None else Decimal("0")


def _string_decimal(value: Decimal | None) -> str:
    return str(quantize(value or Decimal("0"), Decimal("0.01")))


def _dt(value: datetime | None) -> str:
    return value.isoformat() if value else ""


def build_report(
    db: Session,
    year: int,
    month: int,
    template: FinanceSalesReportTemplate | None = None,
) -> dict[str, Any]:
    if not 1 <= month <= 12:
        raise ValueError("非法月份")
    template = template or get_or_create_template(db)
    rules = {**DEFAULT_RULES, **(template.rules or {})}
    start, nxt = month_bounds(year, month)
    orders = (
        db.query(SalesOrder)
        .filter(_valid_order_condition())
        .filter(SalesOrder.ordered_at >= start, SalesOrder.ordered_at < nxt)
        .order_by(SalesOrder.ordered_at, SalesOrder.id)
        .all()
    )
    order_ids = [row.id for row in orders]
    order_nos = [row.order_no for row in orders]
    items_by_order: dict[int, list[SalesOrderItem]] = defaultdict(list)
    if order_ids:
        for item in (
            db.query(SalesOrderItem)
            .filter(SalesOrderItem.order_id.in_(order_ids))
            .order_by(SalesOrderItem.order_id, SalesOrderItem.id)
            .all()
        ):
            items_by_order[item.order_id].append(item)

    refunds_by_order: dict[str, Decimal] = defaultdict(lambda: Decimal("0"))
    if order_nos and rules.get("refund_mode") != "ignore_refund":
        for refund in db.query(AftersalesOrder).filter(AftersalesOrder.order_no.in_(order_nos)).all():
            if refund.refund_amount is not None and _is_valid_refund(refund):
                refunds_by_order[refund.order_no] += refund.refund_amount

    total_quantity = Decimal("0")
    total_sales = Decimal("0")
    total_refunds = Decimal("0")
    rows: list[dict[str, Any]] = []
    platform_summary: dict[str, dict[str, Any]] = {}

    for order in orders:
        platform = order.platform or "未标记平台"
        summary = platform_summary.setdefault(platform, {
            "platform": platform, "orderCount": 0, "quantity": Decimal("0"),
            "salesAmount": Decimal("0"), "refundAmount": Decimal("0"),
        })
        summary["orderCount"] += 1
        refund_amount = refunds_by_order.get(order.order_no, Decimal("0"))
        total_refunds += refund_amount
        summary["refundAmount"] += refund_amount

        items = items_by_order.get(order.id) or [None]
        for index, item in enumerate(items):
            quantity = _money(item.quantity if item else None)
            item_amount = _money(item.amount if item else None)
            discount = _money(item.discount_amount if item else None)
            item_net = item_amount - discount
            total_quantity += quantity
            total_sales += item_net
            summary["quantity"] += quantity
            summary["salesAmount"] += item_net
            first = index == 0
            rows.append({
                "platform": order.platform or "",
                "store_id": str(order.store_id or ""),
                "source_provider": order.source_provider or "",
                "order_no": order.order_no,
                "ordered_at": _dt(order.ordered_at),
                "paid_at": _dt(order.paid_at),
                "order_type": order.order_type or "",
                "order_status": order.order_status or "",
                "pay_status": order.pay_status or "",
                "sku_code": item.sku_code if item else "",
                "goods_name": item.goods_name if item else "",
                "quantity": str(item.quantity) if item and item.quantity is not None else "",
                "unit_price": str(item.unit_price) if item and item.unit_price is not None else "",
                "item_amount": str(item.amount) if item and item.amount is not None else "",
                "discount_amount": str(item.discount_amount) if item and item.discount_amount is not None else "",
                "item_net_amount": str(item_net) if item and (item.amount is not None or item.discount_amount is not None) else "",
                # 订单级值只在第一条明细展示，防止 Excel 用户直接求和时重复。
                "order_amount": str(order.order_amount) if first and order.order_amount is not None else "",
                "paid_amount": str(order.paid_amount) if first and order.paid_amount is not None else "",
                "refund_amount": str(refund_amount) if first and refund_amount else ("0" if first else ""),
                "currency": order.currency or "CNY",
            })

    enabled_fields = [f for f in _normalize_fields(template.fields) if f["enabled"]]
    summary_rows = []
    for item in sorted(platform_summary.values(), key=lambda x: x["salesAmount"], reverse=True):
        summary_rows.append({
            "platform": item["platform"],
            "orderCount": item["orderCount"],
            "quantity": str(item["quantity"]),
            "salesAmount": _string_decimal(item["salesAmount"]),
            "refundAmount": _string_decimal(item["refundAmount"]),
            "netAfterRefund": _string_decimal(item["salesAmount"] - item["refundAmount"]),
        })
    return {
        "year": year,
        "month": month,
        "generatedAt": datetime.now(timezone.utc).isoformat(),
        "template": serialize_template(template),
        "fields": enabled_fields,
        "rules": rules,
        "summary": {
            "orderCount": len(orders),
            "totalQuantity": str(total_quantity),
            "salesAmount": _string_decimal(total_sales),
            "refundAmount": _string_decimal(total_refunds),
            "netAfterRefund": _string_decimal(total_sales - total_refunds),
        },
        "byPlatform": summary_rows,
        "rows": rows,
    }


def _as_excel_value(key: str, raw: Any) -> Any:
    if raw in (None, ""):
        return ""
    field_type = FIELD_REGISTRY.get(key, {}).get("type")
    if field_type in {"money", "number"}:
        try:
            return float(raw)
        except (TypeError, ValueError):
            return raw
    return raw


def to_xlsx(report: dict[str, Any]) -> bytes:
    wb = Workbook()
    summary_ws = wb.active
    summary_ws.title = "销售汇总"
    summary_ws.append(["账期", f"{report['year']}-{report['month']:02d}"])
    summary_ws.append(["订单数", report["summary"]["orderCount"]])
    summary_ws.append(["销售数量", float(report["summary"]["totalQuantity"] or 0)])
    summary_ws.append(["商品净销售额", float(report["summary"]["salesAmount"] or 0)])
    summary_ws.append(["有效退款金额", float(report["summary"]["refundAmount"] or 0)])
    summary_ws.append(["退款后净销售额", float(report["summary"]["netAfterRefund"] or 0)])
    summary_ws.append([])
    summary_ws.append(["平台", "订单数", "销售数量", "商品净销售额", "有效退款金额", "退款后净销售额"])
    for row in report.get("byPlatform", []):
        summary_ws.append([
            row["platform"], row["orderCount"], float(row["quantity"] or 0),
            float(row["salesAmount"] or 0), float(row["refundAmount"] or 0),
            float(row["netAfterRefund"] or 0),
        ])
    summary_ws[8][0].font = Font(bold=True)
    for cell in summary_ws[8]:
        cell.font = Font(bold=True)
    for col in range(4, 7):
        for row in range(2, summary_ws.max_row + 1):
            summary_ws.cell(row=row, column=col).number_format = "0.00"
    summary_ws.freeze_panes = "A8"

    detail_ws = wb.create_sheet("销售明细")
    fields = report.get("fields", [])
    detail_ws.append([field["label"] for field in fields])
    for cell in detail_ws[1]:
        cell.font = Font(bold=True)
    for row in report.get("rows", []):
        detail_ws.append([_as_excel_value(field["key"], row.get(field["key"])) for field in fields])
    detail_ws.freeze_panes = "A2"
    if fields:
        detail_ws.auto_filter.ref = detail_ws.dimensions
        for idx, field in enumerate(fields, start=1):
            field_type = FIELD_REGISTRY.get(field["key"], {}).get("type")
            if field_type == "money":
                for row_idx in range(2, detail_ws.max_row + 1):
                    detail_ws.cell(row=row_idx, column=idx).number_format = "0.00"
            max_len = max(
                len(str(detail_ws.cell(row=r, column=idx).value or ""))
                for r in range(1, min(detail_ws.max_row, 200) + 1)
            )
            detail_ws.column_dimensions[get_column_letter(idx)].width = min(max(max_len + 2, 10), 28)

    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def generated_filename(year: int, month: int) -> str:
    return f"销售汇总_{year}{month:02d}.xlsx"


def has_generated_report(db: Session, company: str, year: int, month: int) -> bool:
    return (
        db.query(ArchiveFile.id)
        .filter_by(
            company=company,
            period_year=year,
            period_month=month,
            category="sales_summary",
            original_name=generated_filename(year, month),
        )
        .first()
        is not None
    )


def generate_and_archive(
    db: Session,
    *,
    company: str,
    year: int,
    month: int,
    actor: str = "system",
    skip_if_exists: bool = False,
) -> dict[str, Any]:
    template = get_or_create_template(db, company)
    if not template.enabled:
        return {"status": "skipped", "reason": "template_disabled"}
    if skip_if_exists and has_generated_report(db, company, year, month):
        return {"status": "exists", "period": f"{year}-{month:02d}"}
    report = build_report(db, year, month, template)
    content = to_xlsx(report)
    archive = finance_service.store_upload(
        db,
        company=company,
        year=year,
        month=month,
        category="sales_summary",
        original_name=generated_filename(year, month),
        content=content,
        actor=actor,
    )
    return {
        "status": "ok",
        "period": f"{year}-{month:02d}",
        "archiveFileId": archive.id,
        "version": archive.version,
        "summary": report["summary"],
    }
