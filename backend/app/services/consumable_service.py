from __future__ import annotations

from datetime import datetime, timezone
from decimal import Decimal

from sqlalchemy import select
from sqlalchemy.orm import Session
from io import BytesIO

from app.models.catalog import ProductSku
from app.models.consumable import Consumable, ConsumableSkuMapping, ConsumableTransaction, InboundConsumableUsage
from app.models.purchase import ExternalPurchaseOrder, PurchaseAllocationItem
from app.utils.money import to_decimal


def _money(value: Decimal | None) -> str | None:
    return f"{to_decimal(value):.4f}" if value is not None else None


def _qty(value: Decimal | None) -> str:
    return f"{to_decimal(value):.4f}"


def _linked_skus(db: Session, consumable_id: int) -> list[dict]:
    rows = db.query(ConsumableSkuMapping, ProductSku).join(
        ProductSku, ProductSku.id == ConsumableSkuMapping.sku_id
    ).filter(ConsumableSkuMapping.consumable_id == consumable_id).order_by(ProductSku.sku_code).all()
    return [{"skuId": sku.id, "skuCode": sku.sku_code, "skuName": sku.sku_name} for _, sku in rows]


def serialize_consumable(row: Consumable, db: Session) -> dict:
    mapping_count = db.query(ConsumableSkuMapping).filter_by(consumable_id=row.id).count()
    available = to_decimal(row.stock_qty) + to_decimal(row.factory_qty)
    min_qty = to_decimal(row.min_stock_qty)
    return {
        "id": row.id,
        "code": row.code,
        "name": row.name,
        "barcode": row.barcode or "",
        "category": row.category,
        "unit": row.unit,
        "purchaseUnitCost": _money(row.purchase_unit_cost),
        "purchasedQty": _qty(row.purchased_qty),
        "usedQty": _qty(row.used_qty),
        "stockQty": _qty(row.stock_qty),
        "factoryQty": _qty(row.factory_qty),
        "transitQty": _qty(row.transit_qty),
        "availableQty": _qty(available),
        "minStockQty": _qty(row.min_stock_qty),
        "taxCode": row.tax_code or "",
        "usageRate": _qty((to_decimal(row.used_qty) / to_decimal(row.purchased_qty)) if to_decimal(row.purchased_qty) > 0 else Decimal("0")),
        "status": row.status,
        "mappingCount": mapping_count,
        "linkedSkus": _linked_skus(db, row.id),
        # 预警口径：可用（自有仓+工厂）≤ 安全库存，或库存为负
        "lowStock": (min_qty > 0 and available <= min_qty) or to_decimal(row.stock_qty) < 0,
    }


def list_consumables(db: Session, search: str = "", status: str | None = None) -> list[dict]:
    query = db.query(Consumable).order_by(Consumable.code, Consumable.id)
    if status:
        query = query.filter(Consumable.status == status)
    term = search.strip()
    if term:
        pattern = f"%{term}%"
        query = query.filter((Consumable.code.ilike(pattern)) | (Consumable.name.ilike(pattern)))
    return [serialize_consumable(row, db) for row in query.limit(500).all()]


def upsert_consumable(
    db: Session,
    *,
    consumable_id: int | None,
    code: str,
    name: str,
    category: str = "",
    unit: str = "个",
    purchase_unit_cost: str | None = None,
    stock_qty: str | None = None,
    min_stock_qty: str | None = None,
    barcode: str | None = None,
    tax_code: str | None = None,
    sku_ids: list[int] | None = None,
) -> Consumable:
    code = code.strip()
    if not code:
        raise ValueError("耗材编码不能为空")
    if stock_qty is not None:
        raise ValueError("库存请通过采购收货、领用或盘点流水调整")
    row = db.get(Consumable, consumable_id) if consumable_id else None
    if row is None:
        row = db.query(Consumable).filter_by(code=code).first()
    if row is None:
        row = Consumable(code=code)
        db.add(row)
    elif row.code != code and db.query(Consumable).filter_by(code=code).first():
        raise ValueError("耗材编码已存在，请换一个编码")
    row.code = code
    row.name = name.strip()
    row.category = category.strip()
    row.unit = unit.strip() or "个"
    if barcode is not None:
        row.barcode = barcode.strip()
    if purchase_unit_cost is not None:
        row.purchase_unit_cost = to_decimal(purchase_unit_cost)
    if min_stock_qty is not None:
        row.min_stock_qty = to_decimal(min_stock_qty)
    if tax_code is not None:
        row.tax_code = tax_code.strip()
    # 关联正品（多对多）：以本次提交为准同步映射；新建映射默认用量 1（第二阶段自动扣减才用到）。
    if sku_ids is not None:
        wanted: dict[int, None] = {}
        for sku_id in sku_ids:
            try:
                wanted[int(sku_id)] = None
            except (TypeError, ValueError):
                raise ValueError("关联正品格式不正确")
        for sku_id in wanted:
            if db.get(ProductSku, sku_id) is None:
                raise ValueError(f"正品 SKU {sku_id} 不存在")
        existing_mappings = {m.sku_id: m for m in db.query(ConsumableSkuMapping).filter_by(consumable_id=row.id).all()}
        for sku_id in wanted:
            if sku_id not in existing_mappings:
                db.add(ConsumableSkuMapping(consumable_id=row.id, sku_id=sku_id, usage_per_unit=Decimal("1"), note="货品档案关联"))
        for sku_id, mapping in existing_mappings.items():
            if sku_id not in wanted:
                db.delete(mapping)
    db.commit()
    db.refresh(row)
    return row


TX_TYPES = {
    "purchase": "采购入库",
    "send_factory": "发往工厂",
    "factory_receive": "工厂收货",
    "consume": "消耗",
    "stocktake": "盘点",
    "loss": "报损",
    "manual": "手工调整",
    # 旧类型别名：adjustment 与 stocktake 同义（有符号差额，记自有仓）
    "adjustment": "盘点",
}
_SIGNED_TX_TYPES = {"stocktake", "adjustment", "manual"}


def record_transaction(
    db: Session,
    *,
    consumable_id: int,
    transaction_type: str,
    quantity: str,
    unit_cost: str | None = None,
    source_type: str = "manual",
    source_id: int | None = None,
    note: str = "",
    request_key: str | None = None,
    commit: bool = True,
    location: str = "own",
) -> ConsumableTransaction:
    row = db.scalar(select(Consumable).where(Consumable.id == consumable_id).with_for_update().execution_options(populate_existing=True))
    if row is None:
        raise ValueError("耗材不存在")
    if transaction_type not in TX_TYPES:
        raise ValueError(f"流水类型必须是 {'、'.join(TX_TYPES)} 之一")
    if location not in {"own", "factory"}:
        raise ValueError("库存位置必须是 own（自有仓）或 factory（工厂）")
    qty = to_decimal(quantity)
    cost = to_decimal(unit_cost) if unit_cost not in (None, "") else None
    if not qty.is_finite() or (cost is not None and (not cost.is_finite() or cost < 0)):
        raise ValueError("数量必须是有效数字，成本不能为负")
    if transaction_type not in _SIGNED_TX_TYPES and qty <= 0:
        raise ValueError("该流水类型的数量必须大于 0")
    if transaction_type in _SIGNED_TX_TYPES and qty == 0:
        raise ValueError("调整数量不能为 0")
    if request_key:
        existing = db.query(ConsumableTransaction).filter_by(request_key=request_key).first()
        if existing:
            if (existing.consumable_id, existing.transaction_type, existing.quantity, existing.unit_cost, existing.note) != (consumable_id, transaction_type, qty, cost, note.strip()):
                raise ValueError("这次登记已提交过不同内容，请刷新后重新操作")
            return existing
    if source_type != "manual" and source_id is not None:
        existing = db.query(ConsumableTransaction).filter_by(
            source_type=source_type, source_id=source_id, consumable_id=consumable_id
        ).first()
        if existing:
            return existing

    stock_before = to_decimal(row.stock_qty)
    factory_before = to_decimal(row.factory_qty)
    transit = to_decimal(row.transit_qty)
    tx_location: str | None = None
    if transaction_type == "purchase":
        # 到货位置：自有仓或工厂（耗材不进吉客云，收货即入本平台对应库存）
        tx_location = location
        if location == "factory":
            row.factory_qty = factory_before + qty
        else:
            row.stock_qty = stock_before + qty
        row.purchased_qty = to_decimal(row.purchased_qty) + qty
    elif transaction_type == "send_factory":
        # 发往工厂：自有仓 → 在途
        row.stock_qty = stock_before - qty
        row.transit_qty = transit + qty
    elif transaction_type == "factory_receive":
        # 工厂收货：在途 → 工厂库存
        if transit - qty < 0:
            raise ValueError(f"在途库存不足（当前在途 {transit}），请先登记发往工厂")
        row.transit_qty = transit - qty
        row.factory_qty = factory_before + qty
    elif transaction_type == "consume":
        # 允许领用后出现负库存：历史采购补录时不能阻断业务，界面会显示低库存，
        # 后续采购入库或盘点调整再补齐；流水仍完整保留。
        row.stock_qty = stock_before - qty
        row.used_qty = to_decimal(row.used_qty) + qty
    elif transaction_type == "loss":
        row.stock_qty = stock_before - qty
    else:  # stocktake / adjustment / manual：有符号差额
        tx_location = location
        if location == "factory":
            row.factory_qty = factory_before + qty
        else:
            row.stock_qty = stock_before + qty
    tx = ConsumableTransaction(
        consumable_id=consumable_id,
        transaction_type=transaction_type,
        quantity=qty,
        unit_cost=cost,
        location=tx_location,
        source_type=source_type,
        source_id=source_id,
        stock_before=stock_before,
        stock_after=to_decimal(row.stock_qty),
        factory_before=factory_before,
        factory_after=to_decimal(row.factory_qty),
        request_key=request_key,
        note=note.strip(),
        occurred_at=datetime.now(timezone.utc),
    )
    db.add(tx)
    if commit:
        db.commit()
        db.refresh(tx)
    else:
        db.flush()
    return tx


def consume_for_purchase_order(db: Session, po: ExternalPurchaseOrder) -> dict[str, str]:
    """采购内容确认后按 SKU 映射自动领用耗材；同一 PO 幂等。"""
    allocations = db.query(PurchaseAllocationItem).filter(
        PurchaseAllocationItem.po_id == po.id,
        PurchaseAllocationItem.sku_id.isnot(None),
        PurchaseAllocationItem.quantity.isnot(None),
    ).all()
    totals: dict[int, Decimal] = {}
    for allocation in allocations:
        for mapping in db.query(ConsumableSkuMapping).filter_by(sku_id=allocation.sku_id).all():
            totals[mapping.consumable_id] = totals.get(mapping.consumable_id, Decimal("0")) + to_decimal(allocation.quantity) * to_decimal(mapping.usage_per_unit)
    consumed = Decimal("0")
    for consumable_id, qty in totals.items():
        record_transaction(
            db, consumable_id=consumable_id, transaction_type="consume", quantity=str(qty),
            source_type="purchase_order", source_id=po.id,
            note=f"采购单 {po.external_order_id} 确认后按 SKU 映射自动领用",
        )
        consumed += qty
    return {"consumableCount": str(len(totals)), "consumedQty": _qty(consumed)}


def inbound_usage_rows(db: Session, link_id: int) -> list[dict]:
    """返回一条入库关联当前确认的耗材使用明细。"""
    rows = db.query(InboundConsumableUsage, Consumable).join(
        Consumable, Consumable.id == InboundConsumableUsage.consumable_id
    ).filter(InboundConsumableUsage.link_id == link_id).order_by(Consumable.code).all()
    return [
        {
            "id": usage.id,
            "consumableId": material.id,
            "consumableCode": material.code,
            "consumableName": material.name,
            "unit": material.unit,
            "quantity": _qty(usage.quantity),
            "note": usage.note,
        }
        for usage, material in rows
    ]


def _reverse_inbound_usage(db: Session, link_id: int) -> None:
    """撤销一条入库关联的耗材扣减，供更换/解除/重新编辑使用。"""
    from app.models.procurement_chain import ProcurementChainLink
    db.scalar(select(ProcurementChainLink).where(ProcurementChainLink.id == link_id).with_for_update())
    usages = db.query(InboundConsumableUsage).filter_by(link_id=link_id).order_by(InboundConsumableUsage.consumable_id).all()
    for usage in usages:
        material = db.scalar(select(Consumable).where(Consumable.id == usage.consumable_id).with_for_update().execution_options(populate_existing=True))
        if material is not None:
            material.stock_qty = to_decimal(material.stock_qty) + to_decimal(usage.quantity)
            material.used_qty = max(Decimal("0"), to_decimal(material.used_qty) - to_decimal(usage.quantity))
        tx = db.query(ConsumableTransaction).filter_by(
            source_type="inbound_link", source_id=link_id, consumable_id=usage.consumable_id
        ).first()
        if tx is not None:
            db.delete(tx)
        db.delete(usage)


def set_inbound_usage(
    db: Session,
    *,
    link_id: int,
    enabled: bool,
    items: list[dict] | None = None,
    note: str = "",
) -> list[dict]:
    """确认一条入库关联是否产生耗材出库，并按明细扣减库存。

    ``enabled=False`` 是明确的“本次不使用耗材”，不是未填写；enabled=True
    必须至少有一条数量大于 0 的耗材明细。重复提交会先冲销旧明细再重算，保持幂等。
    """
    from app.models.procurement_chain import ProcurementChainLink
    from app.models.jackyun import JackyunGoodsDocument

    link = db.scalar(select(ProcurementChainLink).where(ProcurementChainLink.id == link_id).with_for_update().execution_options(populate_existing=True))
    if link is None or link.target_type != "inbound":
        raise ValueError("入库关联不存在")
    if db.get(JackyunGoodsDocument, link.target_id) is None:
        raise ValueError("关联的入库单不存在")
    raw_items = items or []
    if not isinstance(enabled, bool):
        raise ValueError("必须明确选择是否添加耗材使用")
    if not enabled and raw_items:
        raise ValueError("选择不使用耗材时不能填写耗材明细")
    normalized: dict[int, Decimal] = {}
    for item in raw_items:
        try:
            material_id = int(item.get("consumable_id"))
            quantity = to_decimal(item.get("quantity"))
        except (TypeError, ValueError, AttributeError) as exc:
            raise ValueError("耗材明细格式不正确") from exc
        if not quantity.is_finite() or quantity <= 0:
            raise ValueError("耗材使用数量必须大于 0")
        if db.get(Consumable, material_id) is None:
            raise ValueError(f"耗材 {material_id} 不存在")
        normalized[material_id] = normalized.get(material_id, Decimal("0")) + quantity
    if enabled and not normalized:
        raise ValueError("选择使用耗材时至少填写一条耗材及使用数量")

    old_items = db.query(InboundConsumableUsage).filter_by(link_id=link_id).all()
    old_quantities = {item.consumable_id: to_decimal(item.quantity) for item in old_items}
    if link.consumable_usage_decided and link.consumable_usage_enabled == enabled and old_quantities == normalized:
        return inbound_usage_rows(db, link_id)
    all_ids = sorted(set(old_quantities) | set(normalized))
    db.scalars(select(Consumable).where(Consumable.id.in_(all_ids)).order_by(Consumable.id).with_for_update().execution_options(populate_existing=True)).all()
    _reverse_inbound_usage(db, link_id)
    # 先落下旧明细删除，避免同一唯一键的新增先于删除触发冲突。
    db.flush()
    if enabled:
        for material_id, quantity in normalized.items():
            material = db.get(Consumable, material_id)
            assert material is not None
            material.stock_qty = to_decimal(material.stock_qty) - quantity
            material.used_qty = to_decimal(material.used_qty) + quantity
            db.add(InboundConsumableUsage(
                link_id=link_id,
                inbound_document_id=link.target_id,
                consumable_id=material_id,
                quantity=quantity,
                note=note.strip(),
            ))
            db.add(ConsumableTransaction(
                consumable_id=material_id,
                transaction_type="consume",
                quantity=quantity,
                unit_cost=material.purchase_unit_cost,
                source_type="inbound_link",
                source_id=link_id,
                note=note.strip() or f"入库单关联 #{link.target_id} 确认耗材出库",
                occurred_at=datetime.now(timezone.utc),
            ))
    link.consumable_usage_decided = True
    link.consumable_usage_enabled = enabled
    db.commit()
    return inbound_usage_rows(db, link_id)


def list_transactions(db: Session, consumable_id: int, limit: int = 100) -> list[dict]:
    from app.models.consumable_purchase import ConsumableReceipt
    from app.models.procurement_chain import ProcurementChainLink
    rows = db.query(ConsumableTransaction).filter_by(consumable_id=consumable_id).order_by(
        ConsumableTransaction.occurred_at.desc().nullslast(), ConsumableTransaction.id.desc()
    ).limit(limit).all()
    receipt_ids = [row.source_id for row in rows if row.source_type == "consumable_receipt"]
    receipts = {r.id: r for r in db.query(ConsumableReceipt).filter(ConsumableReceipt.id.in_(receipt_ids)).all()} if receipt_ids else {}
    link_ids = [row.source_id for row in rows if row.source_type == "inbound_link"]
    links = {link.id: link for link in db.query(ProcurementChainLink).filter(ProcurementChainLink.id.in_(link_ids)).all()} if link_ids else {}
    return [
        {
            "id": row.id,
            "transactionType": row.transaction_type,
            "quantity": _qty(row.quantity),
            "unitCost": _money(row.unit_cost),
            "location": row.location,
            "stockBefore": _qty(row.stock_before) if row.stock_before is not None else None,
            "stockAfter": _qty(row.stock_after) if row.stock_after is not None else None,
            "factoryBefore": _qty(row.factory_before) if row.factory_before is not None else None,
            "factoryAfter": _qty(row.factory_after) if row.factory_after is not None else None,
            "sourceType": row.source_type,
            "sourceId": row.source_id,
            "note": row.note,
            "occurredAt": row.occurred_at.isoformat() if row.occurred_at else None,
            "purchaseId": receipts[row.source_id].purchase_id if row.source_type == "consumable_receipt" and row.source_id in receipts else None,
            "orderId": links[row.source_id].workbench_order_id if row.source_type == "inbound_link" and row.source_id in links else None,
        }
        for row in rows
    ]


def list_mappings(db: Session, consumable_id: int | None = None, sku_id: int | None = None) -> list[dict]:
    query = db.query(ConsumableSkuMapping, Consumable, ProductSku).join(
        Consumable, Consumable.id == ConsumableSkuMapping.consumable_id
    ).join(ProductSku, ProductSku.id == ConsumableSkuMapping.sku_id)
    if consumable_id is not None:
        query = query.filter(ConsumableSkuMapping.consumable_id == consumable_id)
    if sku_id is not None:
        query = query.filter(ConsumableSkuMapping.sku_id == sku_id)
    return [
        {
            "id": mapping.id,
            "skuId": sku.id,
            "skuCode": sku.sku_code,
            "skuName": sku.sku_name,
            "consumableId": consumable.id,
            "consumableCode": consumable.code,
            "consumableName": consumable.name,
            "usagePerUnit": _qty(mapping.usage_per_unit),
            "note": mapping.note,
        }
        for mapping, consumable, sku in query.order_by(Consumable.code, ProductSku.sku_code).all()
    ]


def upsert_mapping(db: Session, *, mapping_id: int | None, sku_id: int, consumable_id: int, usage_per_unit: str, note: str = "") -> ConsumableSkuMapping:
    if db.get(ProductSku, sku_id) is None or db.get(Consumable, consumable_id) is None:
        raise ValueError("货品或耗材不存在")
    qty = to_decimal(usage_per_unit)
    if qty <= 0:
        raise ValueError("单位消耗量必须大于 0")
    row = db.get(ConsumableSkuMapping, mapping_id) if mapping_id else db.query(ConsumableSkuMapping).filter_by(
        sku_id=sku_id, consumable_id=consumable_id
    ).first()
    if row is None:
        row = ConsumableSkuMapping(sku_id=sku_id, consumable_id=consumable_id)
        db.add(row)
    row.usage_per_unit = qty
    row.note = note.strip()
    db.commit()
    db.refresh(row)
    return row


def remove_mapping(db: Session, mapping_id: int) -> None:
    row = db.get(ConsumableSkuMapping, mapping_id)
    if row is None:
        raise ValueError("耗材映射不存在")
    db.delete(row)
    db.commit()


def import_xlsx(db: Session, content: bytes, filename: str = "consumables.xlsx") -> dict[str, int | str]:
    """读取历史表格的“耗材使用情况”页；只导入可识别的值，不执行原表公式。"""
    if not filename.lower().endswith((".xlsx", ".xlsm")):
        raise ValueError("耗材清单仅支持 .xlsx 或 .xlsm")
    try:
        from openpyxl import load_workbook
        book = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"耗材清单无法读取：{exc}") from exc
    sheet = book["耗材使用情况"] if "耗材使用情况" in book.sheetnames else None
    if sheet is None:
        raise ValueError("未找到“耗材使用情况”工作表")
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise ValueError("耗材工作表为空")
    headers = [str(v).strip() if v is not None else "" for v in rows[0]]
    required = {"条形码", "产品名", "外包装采购数量", "采购单价"}
    if not required.issubset(headers):
        raise ValueError(f"耗材表缺少字段：{', '.join(sorted(required - set(headers)))}")
    idx = {name: headers.index(name) for name in headers if name}
    created = updated = skipped = mappings = 0
    mapping_cache: dict[tuple[int, int], ConsumableSkuMapping] = {}
    def json_value(value):
        if isinstance(value, (datetime,)):
            return value.isoformat()
        if isinstance(value, (Decimal, float, int, str, bool)) or value is None:
            return str(value) if isinstance(value, Decimal) else value
        return str(value)
    for values in rows[1:]:
        code = str(values[idx["条形码"]] or "").strip()
        name = str(values[idx["产品名"]] or "").strip()
        if not code or not name:
            skipped += 1
            continue
        def cell(name: str):
            return values[idx[name]] if name in idx and idx[name] < len(values) else None
        cost = cell("采购单价")
        purchased = cell("外包装采购数量")
        used = cell("使用耗材")
        stock = cell("剩余数量")
        category = str(cell("分类") or "外包装").strip()
        existing = db.query(Consumable).filter_by(code=code).first()
        row = existing or Consumable(code=code)
        row.name = name
        row.category = category
        row.unit = "个"
        if cost is not None and str(cost).strip() not in {"", "None"}:
            row.purchase_unit_cost = to_decimal(str(cost))
        if purchased is not None and isinstance(purchased, (int, float, Decimal)):
            row.purchased_qty = to_decimal(str(purchased))
        if used is not None and isinstance(used, (int, float, Decimal)):
            row.used_qty = to_decimal(str(used))
        if stock is not None and isinstance(stock, (int, float, Decimal)):
            row.stock_qty = to_decimal(str(stock))
        row.raw = {"source": filename, "sheet": "耗材使用情况", "row": {headers[i]: json_value(values[i]) for i in range(min(len(headers), len(values)))}}
        db.add(row)
        if existing: updated += 1
        else: created += 1
    db.flush()
    # “订货”页中条形码、耗材代码、耗材使用量和订购件数-盒形成明确映射；
    # 只有能命中现有吉客云 SKU 的行才自动建立，其他行留给界面人工维护。
    order_sheet = book["订货"] if "订货" in book.sheetnames else None
    if order_sheet is not None:
        order_rows = list(order_sheet.iter_rows(values_only=True))
        if order_rows:
            order_headers = [str(v).strip() if v is not None else "" for v in order_rows[0]]
            oi = {name: order_headers.index(name) for name in order_headers if name}
            for values in order_rows[1:]:
                def oval(name: str):
                    return values[oi[name]] if name in oi and oi[name] < len(values) else None
                sku_code = str(oval("条形码") or "").strip()
                consumable_code = str(oval("耗材代码") or "").strip()
                order_qty = oval("订购件数-盒")
                usage_qty = oval("耗材使用量")
                # 组合装行描述的是一次性拆包关系，不代表单品的固定耗材映射，
                # 不把它推断成单品单位消耗量，避免同一 SKU 出现 1/10 等冲突。
                if "组合装" in str(oval("备注") or ""):
                    continue
                if not sku_code or not consumable_code or order_qty is None or usage_qty is None:
                    continue
                try:
                    ratio = to_decimal(str(usage_qty)) / to_decimal(str(order_qty))
                except (TypeError, ValueError, ArithmeticError):
                    continue
                if ratio <= 0:
                    continue
                sku = db.query(ProductSku).filter((ProductSku.sku_code == sku_code) | (ProductSku.barcode == sku_code)).first()
                material = db.query(Consumable).filter_by(code=consumable_code).first()
                if sku is None or material is None:
                    continue
                key = (sku.id, material.id)
                mapping = mapping_cache.get(key) or db.query(ConsumableSkuMapping).filter_by(sku_id=sku.id, consumable_id=material.id).first()
                if mapping is None:
                    mapping = ConsumableSkuMapping(sku_id=sku.id, consumable_id=material.id, usage_per_unit=ratio, note="由历史订货表自动映射")
                    db.add(mapping)
                    mappings += 1
                elif to_decimal(mapping.usage_per_unit) != ratio:
                    mapping.usage_per_unit = ratio
                    mapping.note = "由历史订货表更新映射"
                    mappings += 1
                mapping_cache[key] = mapping
    db.commit()
    return {"created": created, "updated": updated, "skipped": skipped, "mappings": mappings, "sheet": "耗材使用情况"}
