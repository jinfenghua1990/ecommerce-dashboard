from datetime import datetime
from decimal import Decimal
from io import BytesIO
from uuid import uuid4
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from app.config import settings
from app.models.sales import AftersalesOrder, SalesOrder, SalesOrderItem
from app.services import finance_sales_report_service as svc


def _order_no(prefix: str) -> str:
    return f"{prefix}-{uuid4().hex[:12]}"


def test_default_fields_are_configurable_and_stable():
    fields = svc.default_fields()
    keys = [row["key"] for row in fields]
    enabled = {row["key"] for row in fields if row["enabled"]}
    assert len(keys) == len(set(keys))
    assert "platform" in enabled
    assert "order_no" in enabled
    assert "refund_amount" in enabled
    assert "source_provider" in keys


def test_monthly_sales_report_uses_valid_sales_and_does_not_duplicate_order_values(db_session):
    tz = ZoneInfo(settings.TZ)
    order_no = _order_no("FIN")
    ignored_no = _order_no("UNPAID")

    paid = SalesOrder(
        order_no=order_no, platform="淘宝", order_status="已完成", pay_status="已支付",
        order_amount=Decimal("110"), paid_amount=Decimal("100"),
        ordered_at=datetime(2026, 8, 15, 12, 0, tzinfo=tz),
        paid_at=datetime(2026, 8, 15, 12, 1, tzinfo=tz),
    )
    ignored = SalesOrder(
        order_no=ignored_no, platform="拼多多", order_status="待付款", pay_status="unpaid",
        ordered_at=datetime(2026, 8, 16, 12, 0, tzinfo=tz),
    )
    db_session.add_all([paid, ignored])
    db_session.flush()
    db_session.add_all([
        SalesOrderItem(order_id=paid.id, sku_code="SKU-A", goods_name="商品A",
                       quantity=Decimal("2"), unit_price=Decimal("30"), amount=Decimal("60"), discount_amount=Decimal("10")),
        SalesOrderItem(order_id=paid.id, sku_code="SKU-B", goods_name="商品B",
                       quantity=Decimal("1"), unit_price=Decimal("50"), amount=Decimal("50"), discount_amount=Decimal("0")),
        SalesOrderItem(order_id=ignored.id, sku_code="SKU-X", goods_name="未支付商品",
                       quantity=Decimal("99"), amount=Decimal("999"), discount_amount=Decimal("0")),
        AftersalesOrder(aftersale_no=_order_no("AFTER"), order_no=order_no, type="refund", status="已完成",
                        refund_amount=Decimal("20"), created_at_src=datetime(2026, 8, 20, 10, 0, tzinfo=tz)),
        AftersalesOrder(aftersale_no=_order_no("CANCEL"), order_no=order_no, type="refund", status="已取消",
                        refund_amount=Decimal("500"), created_at_src=datetime(2026, 8, 21, 10, 0, tzinfo=tz)),
    ])
    db_session.flush()

    company = f"pytest-{uuid4().hex}"
    template = svc.get_or_create_template(db_session, company)
    template = svc.save_template(
        db_session, company=company,
        fields=[
            {"key": "order_no", "label": "财务订单号", "enabled": True},
            {"key": "sku_code", "label": "商品编码", "enabled": True},
            {"key": "paid_amount", "label": "实收", "enabled": True},
            {"key": "refund_amount", "label": "退款", "enabled": True},
        ],
        rules={"refund_mode": "recorded_non_cancelled"},
        to_addrs=["finance@example.com"], cc_addrs=[], auto_send=False, send_day=3, send_hour=10,
    )

    report = svc.build_report(db_session, 2026, 8, template)
    assert report["summary"] == {
        "orderCount": 1,
        "totalQuantity": "3.0000",
        "salesAmount": "100.00",
        "refundAmount": "20.00",
        "netAfterRefund": "80.00",
    }
    assert len(report["rows"]) == 2
    assert report["rows"][0]["paid_amount"] == "100"
    assert report["rows"][1]["paid_amount"] == ""
    assert report["rows"][0]["refund_amount"] == "20.0000"
    assert report["rows"][1]["refund_amount"] == ""
    assert all(row["order_no"] == order_no for row in report["rows"])

    workbook = load_workbook(BytesIO(svc.to_xlsx(report)), data_only=True)
    assert workbook.sheetnames == ["销售汇总", "销售明细"]
    headers = [cell.value for cell in workbook["销售明细"][1]]
    assert headers == ["财务订单号", "商品编码", "实收", "退款"]


def test_refund_can_be_disabled_by_template_rule(db_session):
    tz = ZoneInfo(settings.TZ)
    order_no = _order_no("NOREF")
    order = SalesOrder(order_no=order_no, platform="京东", order_status="finished", paid_amount=Decimal("30"),
                       ordered_at=datetime(2026, 7, 10, 12, 0, tzinfo=tz))
    db_session.add(order)
    db_session.flush()
    db_session.add(SalesOrderItem(order_id=order.id, sku_code="SKU-1", goods_name="商品",
                                  quantity=Decimal("1"), amount=Decimal("30"), discount_amount=Decimal("0")))
    db_session.add(AftersalesOrder(aftersale_no=_order_no("REF"), order_no=order_no, type="refund",
                                   status="done", refund_amount=Decimal("10")))
    db_session.flush()

    company = f"pytest-{uuid4().hex}"
    template = svc.get_or_create_template(db_session, company)
    template = svc.save_template(
        db_session, company=company, fields=svc.default_fields(), rules={"refund_mode": "ignore_refund"},
        to_addrs=[], cc_addrs=[], auto_send=False, send_day=3, send_hour=10,
    )
    report = svc.build_report(db_session, 2026, 7, template)
    assert report["summary"]["refundAmount"] == "0.00"
    assert report["summary"]["netAfterRefund"] == "30.00"
